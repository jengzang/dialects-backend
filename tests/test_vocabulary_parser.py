from pathlib import Path

import pytest
from docx import Document

from app.service.vocabulary.location import normalize_location_payload
from app.service.vocabulary.parser import (
    ParsedVocabularyRow,
    parse_bracket_document_text,
    parse_vocabulary_file,
)


def test_parse_reference_xlsx_maps_written_to_standard_word() -> None:
    result = parse_vocabulary_file(Path("tests/息烽.xlsx"), parser_mode="table")

    assert result.parser_mode == "table"
    assert result.errors == []
    assert result.rows[0] == ParsedVocabularyRow(
        standard_word="哥",
        local_expression="哥",
        ipa="ko1",
        notes="",
    )
    assert len(result.rows) > 10


def test_parse_table_accepts_chinese_column_aliases(tmp_path: Path) -> None:
    csv_path = tmp_path / "vocabulary.csv"
    csv_path.write_text(
        "释义,当地讲法,音标,备注\n"
        "太阳,日头,ȵit2 tʰəu2,常用\n",
        encoding="utf-8",
    )

    result = parse_vocabulary_file(csv_path, parser_mode="table")

    assert result.errors == []
    assert result.rows == [
        ParsedVocabularyRow(
            standard_word="太阳",
            local_expression="日头",
            ipa="ȵit2 tʰəu2",
            notes="常用",
        )
    ]


def test_parse_bracket_document_text_extracts_fields() -> None:
    rows, errors = parse_bracket_document_text(
        "太阳（日头）[ȵit2 tʰəu2]{常用}\n"
        "月亮（月光）[ŋye2 kuaŋ1]"
    )

    assert errors == []
    assert rows == [
        ParsedVocabularyRow(
            standard_word="太阳",
            local_expression="日头",
            ipa="ȵit2 tʰəu2",
            notes="常用",
        ),
        ParsedVocabularyRow(
            standard_word="月亮",
            local_expression="月光",
            ipa="ŋye2 kuaŋ1",
            notes="",
        ),
    ]


def test_parse_whitespace_document_text_allows_fields_on_separate_lines() -> None:
    from app.service.vocabulary.parser import parse_whitespace_document_text

    rows, errors = parse_whitespace_document_text(
        "太阳\n"
        "日头\n"
        "ȵit2 tʰəu2\n"
        "常用\n"
        "\n"
        "月亮\n"
        "月光\n"
        "ŋye2 kuaŋ1"
    )

    assert errors == []
    assert rows == [
        ParsedVocabularyRow(
            standard_word="太阳",
            local_expression="日头",
            ipa="ȵit2 tʰəu2",
            notes="常用",
        ),
        ParsedVocabularyRow(
            standard_word="月亮",
            local_expression="月光",
            ipa="ŋye2 kuaŋ1",
            notes="",
        ),
    ]


def test_parse_doc_plain_text_fallback_for_bracket_mode(tmp_path: Path) -> None:
    doc_path = tmp_path / "vocabulary.doc"
    doc_path.write_text("太阳（日头）[ȵit2 tʰəu2]{常用}\n", encoding="utf-8")

    result = parse_vocabulary_file(doc_path, parser_mode="doc_bracket")

    assert result.errors == []
    assert result.rows == [
        ParsedVocabularyRow(
            standard_word="太阳",
            local_expression="日头",
            ipa="ȵit2 tʰəu2",
            notes="常用",
        )
    ]


def test_parse_docx_whitespace_treats_paragraphs_as_records(tmp_path: Path) -> None:
    docx_path = tmp_path / "vocabulary.docx"
    document = Document()
    document.add_paragraph("太阳 日头 ȵit2")
    document.add_paragraph("月亮 月光 ŋye2")
    document.save(docx_path)

    result = parse_vocabulary_file(docx_path, parser_mode="doc_whitespace")

    assert result.errors == []
    assert result.rows == [
        ParsedVocabularyRow(
            standard_word="太阳",
            local_expression="日头",
            ipa="ȵit2",
            notes="",
        ),
        ParsedVocabularyRow(
            standard_word="月亮",
            local_expression="月光",
            ipa="ŋye2",
            notes="",
        ),
    ]


def test_normalize_location_payload_accepts_chinese_aliases() -> None:
    normalized = normalize_location_payload(
        {
            "簡稱": "息烽",
            "經緯度": "106.73,27.10",
            "省": "贵州",
            "市": "贵阳",
            "縣": "息烽",
            "音典分區": "西南",
            "地圖集二分區": "黔中",
        }
    )

    assert normalized.location_name == "息烽"
    assert normalized.coordinates == "106.73,27.10"
    assert normalized.province == "贵州"
    assert normalized.city == "贵阳"
    assert normalized.county == "息烽"
    assert normalized.yindian_region == "西南"
    assert normalized.atlas_region == "黔中"


def test_normalize_location_payload_requires_location_name_and_coordinates() -> None:
    with pytest.raises(ValueError, match="location_name"):
        normalize_location_payload({"coordinates": "106.73,27.10"})

    with pytest.raises(ValueError, match="coordinates"):
        normalize_location_payload({"location_name": "息烽"})
