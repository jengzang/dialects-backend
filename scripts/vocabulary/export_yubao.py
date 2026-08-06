import sqlite3
import os
import re
import json
from collections import defaultdict
from openpyxl import Workbook

YUBao_PATH = os.path.expanduser("~/CodeProject/dialects/dialects-backend/data/yubao.db")
OUTPUT_DIR = os.path.expanduser("~/CodeProject/dialects/dialects-backend/data/yubao_exports")

HEADERS = ["书面词条", "方言讲法", "国际音标", "注释"]

def safe_name(s: str) -> str:
    """Sanitize strings for use in filenames."""
    if not s or s in ("（无）", "(无)"):
        return ""
    return re.sub(r'[/\\:*?"<>|]', '_', s).strip()

def build_location_name(province: str, city: str, county: str, village: str, location: str) -> str:
    """Build a human-readable location name for the filename."""
    parts = []
    for p in [province, city, county, village, location]:
        name = safe_name(p)
        if name:
            parts.append(name)
    full = "".join(parts)
    # Truncate very long location descriptions to avoid filesystem limits
    MAX_LEN = 100
    if len(full) > MAX_LEN:
        # Keep province-city-county, truncate the rest
        short_parts = []
        for p in [province, city, county]:
            name = safe_name(p)
            if name:
                short_parts.append(name)
        short = "".join(short_parts)
        remaining = MAX_LEN - len(short) - 1
        if remaining > 10:
            extra = "".join(safe_name(p) for p in [village, location] if safe_name(p))
            short = short + "_" + extra[:remaining]
        full = short
    return full

def build_dialect_name(lang_cat1: str, lang_cat2: str, lang_cat3: str) -> str:
    """Build dialect name for the filename."""
    parts = []
    for c in [lang_cat1, lang_cat2, lang_cat3]:
        name = safe_name(c)
        if name and name not in ("系属不明",):
            parts.append(name)
    if not parts:
        return "未分类"
    return "-".join(parts)

def _or_empty(v):
    return (v or "") if v is not None else ""

def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    conn = sqlite3.connect(YUBao_PATH)
    conn.row_factory = sqlite3.Row
    cur = conn.cursor()

    # Step 1: Find dominant lang_cat per location
    print("Step 1: Finding dominant dialect classification per location...")
    cur.execute("""
        SELECT province, city, county, village, location,
               lang_cat1, lang_cat2, lang_cat3,
               count(*) as cnt
        FROM vocabulary
        GROUP BY province, city, county, village, location, lang_cat1, lang_cat2, lang_cat3
        ORDER BY cnt DESC
    """)

    location_dialects = defaultdict(lambda: {"max_cnt": 0, "lang_cat1": "", "lang_cat2": "", "lang_cat3": ""})
    for row in cur.fetchall():
        key = (_or_empty(row["province"]), _or_empty(row["city"]), _or_empty(row["county"]),
               _or_empty(row["village"]), _or_empty(row["location"]))
        if row["cnt"] > location_dialects[key]["max_cnt"]:
            location_dialects[key] = {
                "max_cnt": row["cnt"],
                "lang_cat1": row["lang_cat1"] or "",
                "lang_cat2": row["lang_cat2"] or "",
                "lang_cat3": row["lang_cat3"] or "",
            }

    print(f"  Found {len(location_dialects)} distinct locations")

    # Step 2: Export each location
    print("Step 2: Exporting xlsx files...")
    locations_meta = []
    seen_filenames = {}
    success = 0
    skip = 0

    for idx, ((province, city, county, village, location), dialect_info) in enumerate(
        sorted(location_dialects.items(),
               key=lambda x: (x[0][0] or "", x[0][1] or "", x[0][2] or "", x[0][3] or "", x[0][4] or ""))
    ):
        loc_name = build_location_name(province, city, county, village, location)
        dialect_name = build_dialect_name(
            dialect_info["lang_cat1"], dialect_info["lang_cat2"], dialect_info["lang_cat3"]
        )

        if not loc_name:
            loc_name = "未知地点"
        base = f"{loc_name}_{dialect_name}"
        if base in seen_filenames:
            seen_filenames[base] += 1
            base = f"{base}_{seen_filenames[base]}"
        else:
            seen_filenames[base] = 0
        filename = f"{base}.xlsx"
        filepath = os.path.join(OUTPUT_DIR, filename)

        # Build NULL-safe WHERE clause
        cols = ["province", "city", "county", "village", "location"]
        vals = [province, city, county, village, location]
        where_parts = []
        params_coord = []
        params_entries = []
        for c, v in zip(cols, vals):
            if v == "":
                where_parts.append(f"(IFNULL({c},'') = '')")
            else:
                where_parts.append(f"{c} = ?")
                params_coord.append(v)
                params_entries.append(v)

        where_clause = " AND ".join(where_parts)

        # Get coordinates (first non-null lng,lat for this location)
        cur.execute(f"""
            SELECT longitude, latitude FROM vocabulary
            WHERE {where_clause}
              AND longitude IS NOT NULL AND latitude IS NOT NULL
            LIMIT 1
        """, params_coord)
        coord_row = cur.fetchone()
        lng = coord_row["longitude"] if coord_row else None
        lat = coord_row["latitude"] if coord_row else None
        coordinates = f"{lng},{lat}" if lng is not None and lat is not None else ""

        # Fetch all entries for this location
        cur.execute(f"""
            SELECT word, note2, pronunciation, note1
            FROM vocabulary
            WHERE {where_clause}
            ORDER BY word, note2
        """, params_entries)
        rows = cur.fetchall()

        if not rows:
            skip += 1
            continue

        # Write xlsx
        wb = Workbook()
        ws = wb.active
        ws.title = "词汇"
        ws.append(HEADERS)

        for r in rows:
            word = (r["word"] or "").strip()
            local_exp = (r["note2"] or "").strip()
            ipa = (r["pronunciation"] or "").strip()
            notes = (r["note1"] or "").strip()
            ws.append([word, local_exp, ipa, notes])

        # Auto-adjust column widths
        for col_idx, header in enumerate(HEADERS, 1):
            max_width = len(header) * 2  # Chinese chars are wider
            for row_cells in ws.iter_rows(min_col=col_idx, max_col=col_idx, min_row=1, values_only=True):
                for cell_val in row_cells:
                    if cell_val:
                        width = sum(2 if ord(ch) > 127 else 1 for ch in str(cell_val))
                        max_width = max(max_width, width)
            ws.column_dimensions[ws.cell(1, col_idx).column_letter].width = min(max_width + 2, 60)

        wb.save(filepath)
        success += 1

        # Record metadata
        locations_meta.append({
            "filename": filename,
            "location_name": loc_name,
            "coordinates": coordinates,
            "province": safe_name(province),
            "city": safe_name(city),
            "county": safe_name(county),
            "town": safe_name(village),
            "administrative_village": safe_name(location) if safe_name(location) else "",
            "yindian_region": dialect_name,
            "atlas_region": dialect_name,
            "entry_count": len(rows),
            "lang_cat1": dialect_info["lang_cat1"],
            "lang_cat2": dialect_info["lang_cat2"],
            "lang_cat3": dialect_info["lang_cat3"],
        })

        if success % 100 == 0:
            print(f"  ... {success} files done")

    conn.close()

    # Write locations metadata
    meta_path = os.path.join(OUTPUT_DIR, "_locations.json")
    with open(meta_path, "w", encoding="utf-8") as f:
        json.dump(locations_meta, f, ensure_ascii=False, indent=2)

    total_entries = sum(m["entry_count"] for m in locations_meta)
    print(f"\nDone: {success} xlsx files exported, {skip} empty locations skipped")
    print(f"Total entries: {total_entries}")
    print(f"Location metadata: {meta_path}")
    print(f"Output dir: {OUTPUT_DIR}")

if __name__ == "__main__":
    main()
